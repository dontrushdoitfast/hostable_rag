import io
import csv
import logging
from typing import Optional
from bs4 import BeautifulSoup

logger = logging.getLogger(__name__)


class DocumentProcessor:
    """
    Extracts plain text content from various file formats (PDF, DOCX, PPTX, XLSX, TXT, CSV, HTML, Images).
    """

    IMAGE_EXTENSIONS = {"png", "jpg", "jpeg", "tiff", "tif", "bmp", "webp"}

    @staticmethod
    def extract_text(content: bytes, filename: str, content_type: Optional[str] = None) -> str:
        if not content:
            return ""

        ext = filename.split(".")[-1].lower() if "." in filename else ""
        c_type = (content_type or "").lower()

        try:
            if ext == "pdf" or "pdf" in c_type:
                return DocumentProcessor._extract_pdf(content)
            elif ext in ["docx", "doc"] or "wordprocessingml" in c_type:
                return DocumentProcessor._extract_docx(content)
            elif ext in ["pptx", "ppt"] or "presentationml" in c_type:
                return DocumentProcessor._extract_pptx(content)
            elif ext in ["xlsx", "xls"] or "spreadsheetml" in c_type:
                return DocumentProcessor._extract_xlsx(content)
            elif ext in ["html", "htm"] or "html" in c_type:
                return DocumentProcessor._extract_html(content)
            elif ext == "csv" or "csv" in c_type:
                return DocumentProcessor._extract_csv(content)
            elif ext in DocumentProcessor.IMAGE_EXTENSIONS or "image/" in c_type:
                return DocumentProcessor._extract_image(content, filename)
            else:
                # Default text decoding fallback
                return content.decode("utf-8", errors="replace")
        except Exception as e:
            logger.error(f"Error parsing document {filename}: {e}")
            if ext in DocumentProcessor.IMAGE_EXTENSIONS or "image/" in c_type:
                return f"[Image file: {filename}]"
            return content.decode("utf-8", errors="replace")

    @staticmethod
    def _extract_pdf(content: bytes) -> str:
        import pypdf

        reader = pypdf.PdfReader(io.BytesIO(content))
        text_parts = []
        for page in reader.pages:
            t = page.extract_text()
            if t:
                text_parts.append(t)
        return "\n\n".join(text_parts)

    @staticmethod
    def _extract_docx(content: bytes) -> str:
        import docx

        doc = docx.Document(io.BytesIO(content))
        text_parts = [p.text for p in doc.paragraphs if p.text]
        for table in doc.tables:
            for row in table.rows:
                row_str = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if row_str:
                    text_parts.append(row_str)
        return "\n".join(text_parts)

    @staticmethod
    def _extract_pptx(content: bytes) -> str:
        import pptx

        prs = pptx.Presentation(io.BytesIO(content))
        text_parts = []
        for slide in prs.slides:
            for shape in slide.shapes:
                if shape.has_text_frame:
                    for paragraph in shape.text_frame.paragraphs:
                        if paragraph.text:
                            text_parts.append(paragraph.text)
                if shape.has_table:
                    for row in shape.table.rows:
                        row_str = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                        if row_str:
                            text_parts.append(row_str)
        return "\n".join(text_parts)

    @staticmethod
    def _extract_xlsx(content: bytes) -> str:
        import openpyxl

        wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
        text_parts = []
        for sheet in wb.worksheets:
            text_parts.append(f"Sheet: {sheet.title}")
            for row in sheet.iter_rows(values_only=True):
                row_str = " | ".join([str(val) for val in row if val is not None])
                if row_str:
                    text_parts.append(row_str)
        return "\n".join(text_parts)

    @staticmethod
    def _extract_html(content: bytes) -> str:
        soup = BeautifulSoup(content, "html.parser")
        return soup.get_text(separator="\n", strip=True)

    @staticmethod
    def _extract_csv(content: bytes) -> str:
        text = content.decode("utf-8", errors="replace")
        reader = csv.reader(io.StringIO(text))
        rows = [" | ".join(row) for row in reader if row]
        return "\n".join(rows)

    @staticmethod
    def _extract_image(content: bytes, filename: str) -> str:
        from PIL import Image

        try:
            image = Image.open(io.BytesIO(content))
            width, height = image.size
        except Exception:
            return f"[Image file: {filename}]"

        # Attempt OCR if pytesseract is available
        try:
            import pytesseract

            ocr_text = pytesseract.image_to_string(image).strip()
            if ocr_text:
                return f"[Image OCR - {filename}]\n{ocr_text}"
        except Exception:
            pass

        return f"[Image file: {filename} ({width}x{height})]"
