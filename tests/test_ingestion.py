import io
import pytest
from pathlib import Path
from PIL import Image
from docx import Document
from pptx import Presentation
import openpyxl

from app.ingestion.processor import DocumentProcessor
from app.ingestion.sync_engine import SyncEngine
from app.connectors.dummy_connector import DummySharePointConnector


def test_document_processor_empty_and_text():
    assert DocumentProcessor.extract_text(b"", "empty.txt") == ""
    text = DocumentProcessor.extract_text(b"Hello world!", "test.txt", "text/plain")
    assert text == "Hello world!"


def test_document_processor_html():
    html_content = b"<html><body><h1>Title</h1><p>Body content here.</p></body></html>"
    text = DocumentProcessor.extract_text(html_content, "index.html", "text/html")
    assert "Title" in text
    assert "Body content here." in text


def test_document_processor_csv():
    csv_content = b"header1,header2\nval1,val2"
    text = DocumentProcessor.extract_text(csv_content, "data.csv", "text/csv")
    assert "header1 | header2" in text
    assert "val1 | val2" in text


def test_document_processor_image():
    # Generate a simple 10x10 RGB image
    img = Image.new("RGB", (10, 10), color="blue")
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    content = buf.getvalue()

    text = DocumentProcessor.extract_text(content, "diagram.png", "image/png")
    assert "diagram.png" in text
    assert "10x10" in text or "diagram.png" in text


def test_document_processor_docx_with_table():
    doc = Document()
    doc.add_paragraph("Introduction paragraph.")
    table = doc.add_table(rows=1, cols=2)
    row_cells = table.rows[0].cells
    row_cells[0].text = "Key"
    row_cells[1].text = "Value"

    buf = io.BytesIO()
    doc.save(buf)
    content = buf.getvalue()

    text = DocumentProcessor.extract_text(content, "document.docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document")
    assert "Introduction paragraph." in text
    assert "Key | Value" in text


def test_document_processor_pptx_with_table():
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank layout
    shape = slide.shapes.add_table(1, 2, 0, 0, 100, 100)
    shape.table.cell(0, 0).text = "Slide Col 1"
    shape.table.cell(0, 1).text = "Slide Col 2"

    buf = io.BytesIO()
    prs.save(buf)
    content = buf.getvalue()

    text = DocumentProcessor.extract_text(content, "presentation.pptx", "application/vnd.openxmlformats-officedocument.presentationml.presentation")
    assert "Slide Col 1 | Slide Col 2" in text


def test_document_processor_xlsx():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws.append(["ColA", "ColB"])
    ws.append(["ValA", "ValB"])

    buf = io.BytesIO()
    wb.save(buf)
    content = buf.getvalue()

    text = DocumentProcessor.extract_text(content, "data.xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    assert "Sheet: Summary" in text
    assert "ColA | ColB" in text
    assert "ValA | ValB" in text


@pytest.mark.asyncio
async def test_sync_engine_run_and_lifecycle(tmp_path: Path):
    folder = tmp_path / "Docs"
    folder.mkdir()
    (folder / "doc1.txt").write_text("Text content for document 1.")

    connector = DummySharePointConnector(root_dir=str(tmp_path), auto_seed=False)
    engine = SyncEngine(connector=connector)

    stats = await engine.run_sync()
    assert stats["status"] == "success"
    assert stats["documents_added_or_updated"] == 1
    assert stats["documents_deleted"] == 0

    # Test starting and cleanly stopping background sync
    engine.start_scheduled_sync()
    assert engine._bg_task is not None
    engine.stop_scheduled_sync()
    assert engine._bg_task is None
